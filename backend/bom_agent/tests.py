import io
from datetime import timedelta
from urllib.parse import parse_qs, urlparse
from unittest.mock import patch
import os

from django.contrib.auth.models import User
from django.test import TestCase
from django.core import signing
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from bom_agent.excel_export import render_bom_workbook, upload_bom_workbook
from bom_agent.gmail_integration import decrypt_secret, upsert_gmail_credential
from bom_agent.models import BomLineItem, BomResearchRun, SupplierQuote, TeamContact, TeamRequest, GmailCredential
from bom_agent.research_pipeline import run_bom_research
from projects.models import Project


class BomWorkbookTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(
            username="bom-user",
            email="bom@example.com",
            password="x",
        )
        self.project = Project.objects.create(name="Orbital Valve", owner=self.user)
        self.run = BomResearchRun.objects.create(
            project=self.project,
            created_by=self.user,
            status="researching",
            inputs_json={
                "production_volume": "50–100 units",
                "timeline": "8 weeks",
                "budget": "$125.50",
                "exchange_rates": "1.12",
                "shipping_assumptions": "1.18",
                "compliance": ["RoHS", "REACH"],
            },
            completed_at=timezone.now(),
        )
        self.item = BomLineItem.objects.create(
            run=self.run,
            part_name="Valve Body",
            part_number="VB-001",
            material_spec="6061-T6 Aluminum",
            quantity=12,
            status="sourced",
        )
        self.run.research_log = [
            {
                "ts": "10:00:00",
                "type": "search",
                "message": "Search query: valve body machining supplier",
                "item_id": self.item.id,
                "part_name": self.item.part_name,
                "query": "valve body machining supplier",
            }
        ]
        self.run.results_json = {
            f"item_{self.item.id}": {
                "material": "6061-T6 Aluminum",
                "risk_flags": ["single-source", "geographic-concentration"],
                "risk_summary": "Only one AVL supplier was found in the required region.",
                "summary": "Preferred supplier is Apex Machining based on cost and lead time.",
            }
        }
        self.run.save(update_fields=["research_log", "results_json"])

        SupplierQuote.objects.create(
            line_item=self.item,
            supplier_name="Apex Machining",
            unit_price="8.50",
            moq=10,
            lead_time_days=21,
            tooling_cost="250.00",
            landed_cost_usd="9.10",
            source_url="https://example.com/apex",
            notes="AVL supplier [Country: USA]",
            is_avl=True,
        )
        SupplierQuote.objects.create(
            line_item=self.item,
            supplier_name="Global Fabrication",
            unit_price="9.40",
            moq=25,
            lead_time_days=35,
            tooling_cost="180.00",
            landed_cost_usd="10.00",
            source_url="https://example.com/global",
            notes="Secondary quote [Country: China]",
            is_avl=False,
        )

    def test_render_bom_workbook_has_required_sheets_formulas_and_formatting(self):
        workbook_bytes = render_bom_workbook(self.run)
        workbook = load_workbook(io.BytesIO(workbook_bytes), data_only=False)
        try:
            self.assertEqual(
                workbook.sheetnames,
                [
                    "BOM Summary",
                    "Supplier Comparison",
                    "Assumptions & Inputs",
                    "Research Log",
                    "Risk Matrix",
                ],
            )

            summary = workbook["BOM Summary"]
            self.assertEqual(summary["G2"].value, "=E2*F2")
            self.assertEqual(summary["G3"].value, "=SUM(G2:G2)")
            self.assertEqual(summary.freeze_panes, "A2")

            comparison = workbook["Supplier Comparison"]
            self.assertEqual(len(comparison.conditional_formatting), 2)
            self.assertEqual(comparison.freeze_panes, "A2")

            inputs = workbook["Assumptions & Inputs"]
            self.assertEqual(inputs["C8"].font.color.rgb, "FF0000FF")
            self.assertTrue(inputs["D12"].value.startswith("="))
            self.assertEqual(inputs["D12"].font.color.rgb, "FF000000")

            research = workbook["Research Log"]
            self.assertEqual(research["E4"].value, "https://example.com/apex")

            risk = workbook["Risk Matrix"]
            self.assertEqual(risk["C2"].value, "Flagged")
        finally:
            workbook.close()

    def test_upload_bom_workbook_uses_expected_s3_key(self):
        expected_key = (
            f"bom-reports/{self.run.project_id}/run-{self.run.id}/"
            f"bom-research-run-{self.run.id}.xlsx"
        )
        with patch("bom_agent.excel_export.upload_bytes_to_s3") as upload_mock:
            key = upload_bom_workbook(self.run)

        self.assertEqual(key, expected_key)
        upload_mock.assert_called_once()
        self.assertEqual(
            upload_mock.call_args.kwargs["content_type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )


class BomRunExcelViewTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(
            username="view-user",
            email="view@example.com",
            password="x",
        )
        self.project = Project.objects.create(name="Excel Project", owner=self.user)
        self.run = BomResearchRun.objects.create(
            project=self.project,
            created_by=self.user,
        )
        self.client = APIClient()
        self.client.force_authenticate(user=self.user)
        upsert_gmail_credential(
            user=self.user,
            gmail_address="owner@gmail.com",
            access_token="access-token",
            refresh_token="refresh-token",
            token_expires_at=timezone.now() + timedelta(hours=1),
            scopes_json=["gmail.send", "gmail.readonly"],
        )

    def test_bom_run_excel_view_returns_404_without_excel(self):
        response = self.client.get(f"/api/bom/runs/{self.run.id}/excel/")
        self.assertEqual(response.status_code, 404)

    def test_bom_run_excel_view_returns_presigned_url(self):
        self.run.excel_s3_key = "bom-reports/1/run-1/report.xlsx"
        self.run.save(update_fields=["excel_s3_key"])

        with patch(
            "files_api.s3_service.generate_presigned_url",
            return_value="https://example.com/download.xlsx?sig=1",
        ) as presign_mock, patch(
            "files_api.s3_service.get_s3_object_metadata",
            return_value={
                "size": 4096,
                "last_modified": timezone.now(),
                "content_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            },
        ) as metadata_mock:
            response = self.client.get(f"/api/bom/runs/{self.run.id}/excel/")

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload["url"], "https://example.com/download.xlsx?sig=1")
        self.assertEqual(payload["file_size"], 4096)
        self.assertEqual(payload["filename"], "report.xlsx")
        self.assertIsNotNone(payload["generated_at"])
        presign_mock.assert_called_once_with(
            "bom-reports/1/run-1/report.xlsx",
            expires=3600,
        )
        metadata_mock.assert_called_once_with("bom-reports/1/run-1/report.xlsx")


class BomRunListViewTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(
            username="history-user",
            email="history@example.com",
            password="x",
        )
        self.other_user = User.objects.create_user(
            username="history-other",
            email="history-other@example.com",
            password="x",
        )
        self.project = Project.objects.create(name="History Project", owner=self.user)
        self.other_project = Project.objects.create(name="Other Project", owner=self.other_user)
        self.run_one = BomResearchRun.objects.create(
            project=self.project,
            created_by=self.user,
            status="completed",
            completed_at=timezone.now(),
        )
        BomLineItem.objects.create(run=self.run_one, part_name="Bracket", quantity=4, status="sourced")
        self.run_two = BomResearchRun.objects.create(
            project=self.project,
            created_by=self.user,
            status="researching",
        )
        BomLineItem.objects.create(run=self.run_two, part_name="Valve", quantity=2, status="researching")
        other_run = BomResearchRun.objects.create(
            project=self.other_project,
            created_by=self.other_user,
            status="completed",
        )
        BomLineItem.objects.create(run=other_run, part_name="Hidden", quantity=1, status="sourced")
        self.client = APIClient()
        self.client.force_authenticate(user=self.user)

    def test_bom_run_list_filters_by_project_and_user(self):
        response = self.client.get(f"/api/bom/runs/?project_id={self.project.id}")
        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(len(payload), 2)
        self.assertEqual(payload[0]["id"], self.run_two.id)
        self.assertEqual(payload[1]["id"], self.run_one.id)
        self.assertEqual(len(payload[0]["line_items"]), 1)


class BomResearchPipelineTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(
            username="pipeline-user",
            email="pipeline@example.com",
            password="x",
        )
        self.project = Project.objects.create(name="Pipeline Project", owner=self.user)

    def test_run_bom_research_generates_excel_key(self):
        run = BomResearchRun.objects.create(
            project=self.project,
            created_by=self.user,
            status="gathering_inputs",
            inputs_json={"production_volume": "1–10 units"},
            completed_at=timezone.now() - timedelta(hours=1),
        )
        item = BomLineItem.objects.create(
            run=run,
            part_name="Bracket",
            quantity=4,
            status="pending",
        )

        def stub_research_item(*args, **kwargs):
            BomLineItem.objects.filter(pk=item.pk).update(status="sourced")

        expected_key = (
            f"bom-reports/{self.project.id}/run-{run.id}/"
            f"bom-research-run-{run.id}.xlsx"
        )
        with patch("assistant.rag.find_library_docs", return_value=[]), patch(
            "bom_agent.research_pipeline._research_item", side_effect=stub_research_item
        ), patch(
            "bom_agent.excel_export.upload_bom_workbook",
            return_value=expected_key,
        ) as upload_mock:
            run_bom_research(run.id)

        run.refresh_from_db()
        self.assertEqual(run.status, "completed")
        self.assertEqual(run.excel_s3_key, expected_key)
        self.assertIsNotNone(run.completed_at)
        upload_mock.assert_called_once()


class TeamRequestApiTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(
            username="team-request-user",
            email="team-request@example.com",
            password="x",
        )
        self.project = Project.objects.create(name="Atlas Project", owner=self.user)
        self.run = BomResearchRun.objects.create(
            project=self.project,
            created_by=self.user,
            status="awaiting_team_input",
        )
        self.item = BomLineItem.objects.create(
            run=self.run,
            part_name="Seal Plate",
            quantity=3,
            material_spec="",
            status="needs_input",
        )
        self.material_contact = TeamContact.objects.create(
            workspace_owner=self.user,
            added_by=self.user,
            full_name="Mina Materials",
            email="mina@example.com",
            role="Materials Engineer",
            department="Engineering",
            expertise_tags=["materials"],
            preferred_channel="email",
        )
        self.procurement_contact = TeamContact.objects.create(
            workspace_owner=self.user,
            added_by=self.user,
            full_name="Priya Procurement",
            email="priya@example.com",
            role="Procurement Lead",
            department="Supply Chain",
            expertise_tags=["procurement", "program"],
            preferred_channel="email",
        )
        self.client = APIClient()
        self.client.force_authenticate(user=self.user)

    def test_draft_emails_generates_requests_for_unanswered_questions(self):
        response = self.client.post(f"/api/bom/runs/{self.run.id}/draft-emails/")
        self.assertEqual(response.status_code, 201)
        payload = response.json()
        self.assertEqual(payload["created_count"], 5)
        self.assertEqual(TeamRequest.objects.filter(run=self.run).count(), 5)
        first_request = TeamRequest.objects.filter(run=self.run).order_by("created_at").first()
        self.assertIn(self.project.name, first_request.email_subject)
        self.assertTrue(first_request.email_body)
        self.run.refresh_from_db()
        self.assertEqual(self.run.status, "awaiting_team_input")

    def test_draft_can_be_edited_but_cannot_send_until_approved(self):
        request_obj = TeamRequest.objects.create(
            run=self.run,
            contact=self.procurement_contact,
            recipient_email=self.procurement_contact.email,
            recipient_name=self.procurement_contact.full_name,
            question="What is the target unit cost ceiling?",
            channel="email",
            email_subject="Initial draft",
            email_body="Please advise.",
            status="draft",
        )

        patch_response = self.client.patch(
            f"/api/bom/runs/{self.run.id}/emails/{request_obj.id}/",
            {
                "recipient_email": "new-recipient@example.com",
                "recipient_name": "New Recipient",
                "email_subject": "Edited draft",
                "email_body": "Updated body",
            },
            format="json",
        )
        self.assertEqual(patch_response.status_code, 200)
        request_obj.refresh_from_db()
        self.assertEqual(request_obj.recipient_email, "new-recipient@example.com")
        self.assertEqual(request_obj.status, "draft")

        send_response = self.client.post(
            f"/api/bom/runs/{self.run.id}/emails/{request_obj.id}/send/"
        )
        self.assertEqual(send_response.status_code, 400)
        self.assertIn("approved", send_response.json()["detail"])

        approve_response = self.client.post(
            f"/api/bom/runs/{self.run.id}/emails/{request_obj.id}/approve/"
        )
        self.assertEqual(approve_response.status_code, 200)
        request_obj.refresh_from_db()
        self.assertEqual(request_obj.status, "approved")
        self.assertIsNotNone(request_obj.approved_at)

    def test_draft_can_be_discarded(self):
        request_obj = TeamRequest.objects.create(
            run=self.run,
            contact=self.procurement_contact,
            recipient_email=self.procurement_contact.email,
            recipient_name=self.procurement_contact.full_name,
            question="What is the target unit cost ceiling?",
            question_key="budget",
            channel="email",
            email_subject="Initial draft",
            email_body="Please advise.",
            status="draft",
        )

        response = self.client.delete(
            f"/api/bom/runs/{self.run.id}/emails/{request_obj.id}/"
        )
        self.assertEqual(response.status_code, 204)
        self.assertFalse(TeamRequest.objects.filter(pk=request_obj.id).exists())

    def test_send_and_follow_up_require_approval_and_cap_follow_up(self):
        request_obj = TeamRequest.objects.create(
            run=self.run,
            contact=self.material_contact,
            line_item=self.item,
            recipient_email=self.material_contact.email,
            recipient_name=self.material_contact.full_name,
            question="No material spec found for Seal Plate. Which material applies?",
            question_key=f"material_file_{self.item.id}",
            channel="email",
            email_subject="Material question",
            email_body="Please confirm the material.",
            status="draft",
        )

        self.client.post(f"/api/bom/runs/{self.run.id}/emails/{request_obj.id}/approve/")
        with patch(
            "bom_agent.bom_views.send_team_request_via_gmail",
            return_value={"gmail_address": "owner@gmail.com", "message_id": "msg-1", "thread_id": "thread-1"},
        ):
            send_response = self.client.post(
                f"/api/bom/runs/{self.run.id}/emails/{request_obj.id}/send/"
            )
        self.assertEqual(send_response.status_code, 200)
        request_obj.refresh_from_db()
        self.assertEqual(request_obj.status, "sent")
        self.assertIsNotNone(request_obj.sent_at)
        self.assertEqual(request_obj.gmail_message_id, "msg-1")
        self.assertEqual(request_obj.gmail_thread_id, "thread-1")

        follow_up_response = self.client.post(
            f"/api/bom/runs/{self.run.id}/emails/{request_obj.id}/follow-up/"
        )
        self.assertEqual(follow_up_response.status_code, 200)
        request_obj.refresh_from_db()
        self.assertEqual(request_obj.status, "draft")
        self.assertIsNone(request_obj.approved_at)
        self.assertTrue(request_obj.email_subject.startswith("Follow-up:"))

        send_without_approval = self.client.post(
            f"/api/bom/runs/{self.run.id}/emails/{request_obj.id}/send/"
        )
        self.assertEqual(send_without_approval.status_code, 400)

        self.client.post(f"/api/bom/runs/{self.run.id}/emails/{request_obj.id}/approve/")
        with patch(
            "bom_agent.bom_views.send_team_request_via_gmail",
            return_value={"gmail_address": "owner@gmail.com", "message_id": "msg-2", "thread_id": "thread-1"},
        ):
            follow_up_send_response = self.client.post(
                f"/api/bom/runs/{self.run.id}/emails/{request_obj.id}/send/"
            )
        self.assertEqual(follow_up_send_response.status_code, 200)
        request_obj.refresh_from_db()
        self.assertEqual(request_obj.status, "follow_up")
        self.assertEqual(request_obj.follow_up_count, 1)

        second_follow_up = self.client.post(
            f"/api/bom/runs/{self.run.id}/emails/{request_obj.id}/follow-up/"
        )
        self.assertEqual(second_follow_up.status_code, 400)
        self.assertIn("Only one auto follow-up", second_follow_up.json()["detail"])

    def test_approve_all_and_send_all_processes_all_drafts(self):
        request_one = TeamRequest.objects.create(
            run=self.run,
            contact=self.procurement_contact,
            recipient_email=self.procurement_contact.email,
            recipient_name=self.procurement_contact.full_name,
            question="What is the target unit cost ceiling?",
            channel="email",
            email_subject="Budget question",
            email_body="Please confirm budget.",
            status="draft",
        )
        request_two = TeamRequest.objects.create(
            run=self.run,
            contact=self.material_contact,
            recipient_email=self.material_contact.email,
            recipient_name=self.material_contact.full_name,
            question="No material spec found for Seal Plate. Which material applies?",
            channel="email",
            email_subject="Material question",
            email_body="Please confirm material.",
            status="draft",
        )

        approve_all = self.client.post(f"/api/bom/runs/{self.run.id}/emails/approve-all/")
        self.assertEqual(approve_all.status_code, 200)
        self.assertEqual(approve_all.json()["approved_count"], 2)
        request_one.refresh_from_db()
        request_two.refresh_from_db()
        self.assertIsNotNone(request_one.approved_at)
        self.assertIsNotNone(request_two.approved_at)

        with patch(
            "bom_agent.bom_views.send_team_request_via_gmail",
            side_effect=[
                {"gmail_address": "owner@gmail.com", "message_id": "msg-a", "thread_id": "thread-a"},
                {"gmail_address": "owner@gmail.com", "message_id": "msg-b", "thread_id": "thread-b"},
            ],
        ):
            send_all = self.client.post(f"/api/bom/runs/{self.run.id}/emails/send-all/")
        self.assertEqual(send_all.status_code, 200)
        request_one.refresh_from_db()
        request_two.refresh_from_db()
        self.assertEqual(request_one.status, "sent")
        self.assertEqual(request_two.status, "sent")

    def test_overdue_requests_are_flagged_in_email_list(self):
        request_obj = TeamRequest.objects.create(
            run=self.run,
            contact=self.procurement_contact,
            recipient_email=self.procurement_contact.email,
            recipient_name=self.procurement_contact.full_name,
            question="What is the target unit cost ceiling?",
            question_key="budget",
            channel="email",
            email_subject="Budget question",
            email_body="Please confirm budget.",
            status="sent",
            sent_at=timezone.now() - timedelta(hours=49),
        )
        response = self.client.get(f"/api/bom/runs/{self.run.id}/emails/")
        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload[0]["id"], request_obj.id)
        self.assertTrue(payload[0]["is_overdue"])


class GmailIntegrationTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(
            username="gmail-user",
            email="gmail-user@example.com",
            password="x",
        )
        self.project = Project.objects.create(name="Reply Project", owner=self.user)
        self.run = BomResearchRun.objects.create(
            project=self.project,
            created_by=self.user,
            status="awaiting_team_input",
        )
        self.item = BomLineItem.objects.create(
            run=self.run,
            part_name="housing_top",
            quantity=2,
            material_spec="",
            status="needs_input",
        )
        self.contact = TeamContact.objects.create(
            workspace_owner=self.user,
            added_by=self.user,
            full_name="Sarah Stone",
            email="sarah@example.com",
            role="Materials Engineer",
            department="Engineering",
            expertise_tags=["materials"],
            preferred_channel="email",
        )
        self.client = APIClient()
        self.client.force_authenticate(user=self.user)

    def test_gmail_credential_endpoint_encrypts_tokens(self):
        response = self.client.patch(
            "/api/team/gmail/credential/",
            {
                "gmail_address": "owner@gmail.com",
                "access_token": "plain-access",
                "refresh_token": "plain-refresh",
                "token_expires_at": (timezone.now() + timedelta(hours=1)).isoformat(),
                "scopes_json": ["gmail.send", "gmail.readonly"],
            },
            format="json",
        )
        self.assertEqual(response.status_code, 200)
        credential = GmailCredential.objects.get(user=self.user)
        self.assertNotEqual(credential.encrypted_access_token, "plain-access")
        self.assertNotEqual(credential.encrypted_refresh_token, "plain-refresh")
        self.assertEqual(decrypt_secret(credential.encrypted_access_token), "plain-access")
        self.assertEqual(decrypt_secret(credential.encrypted_refresh_token), "plain-refresh")

    @patch.dict(
        os.environ,
        {
            "GOOGLE_OAUTH_CLIENT_ID": "test-client-id",
            "GOOGLE_OAUTH_CLIENT_SECRET": "test-client-secret",
        },
        clear=False,
    )
    def test_gmail_oauth_start_returns_authorization_url(self):
        response = self.client.post(
            "/api/team/gmail/oauth/start/",
            {"next_url": "http://localhost:5173/settings"},
            format="json",
        )

        self.assertEqual(response.status_code, 200)
        auth_url = response.json()["auth_url"]
        parsed = urlparse(auth_url)
        self.assertEqual(parsed.netloc, "accounts.google.com")
        query = parse_qs(parsed.query)
        self.assertEqual(query["redirect_uri"][0], "http://testserver/api/team/gmail/oauth/callback/")
        self.assertEqual(query["response_type"][0], "code")
        state = signing.loads(query["state"][0], salt="bom-agent-gmail-oauth", max_age=600)
        self.assertEqual(state["user_id"], self.user.id)
        self.assertEqual(state["next_url"], "http://localhost:5173/settings")

    @patch("bom_agent.views.fetch_gmail_profile", return_value={"emailAddress": "owner@gmail.com"})
    @patch(
        "bom_agent.views.exchange_google_oauth_code",
        return_value={
            "access_token": "oauth-access",
            "refresh_token": "oauth-refresh",
            "expires_in": 3600,
            "scope": "https://www.googleapis.com/auth/gmail.send https://www.googleapis.com/auth/gmail.readonly",
        },
    )
    def test_gmail_oauth_callback_persists_credential_and_redirects(
        self,
        _exchange_mock,
        _profile_mock,
    ):
        state = signing.dumps(
            {"user_id": self.user.id, "next_url": "http://localhost:5173/settings"},
            salt="bom-agent-gmail-oauth",
        )

        response = self.client.get(
            f"/api/team/gmail/oauth/callback/?code=oauth-code&state={state}"
        )

        self.assertEqual(response.status_code, 302)
        self.assertEqual(
            response["Location"],
            "http://localhost:5173/settings?gmail_oauth=connected",
        )
        credential = GmailCredential.objects.get(user=self.user)
        self.assertEqual(credential.gmail_address, "owner@gmail.com")
        self.assertEqual(decrypt_secret(credential.encrypted_access_token), "oauth-access")
        self.assertEqual(decrypt_secret(credential.encrypted_refresh_token), "oauth-refresh")
        self.assertIn("https://www.googleapis.com/auth/gmail.send", credential.scopes_json)

    def test_poll_endpoint_marks_answered_and_updates_line_item(self):
        upsert_gmail_credential(
            user=self.user,
            gmail_address="owner@gmail.com",
            access_token="access-token",
            refresh_token="refresh-token",
            token_expires_at=timezone.now() + timedelta(hours=1),
            scopes_json=["gmail.send", "gmail.readonly"],
        )
        request_obj = TeamRequest.objects.create(
            run=self.run,
            contact=self.contact,
            line_item=self.item,
            recipient_email=self.contact.email,
            recipient_name=self.contact.full_name,
            question="No material spec found for housing_top. Which material applies?",
            question_key=f"material_file_{self.item.id}",
            channel="email",
            email_subject="Material question",
            email_body="Please confirm material.",
            status="sent",
            gmail_message_id="msg-1",
            gmail_thread_id="thread-1",
            sent_at=timezone.now() - timedelta(hours=2),
        )

        inbound_message = {
            "id": "reply-1",
            "payload": {
                "headers": [{"name": "From", "value": "Sarah Stone <sarah@example.com>"}],
                "body": {},
                "parts": [
                    {
                        "mimeType": "text/plain",
                        "body": {
                            "data": "NjA2MS1UNiBBbHVtaW51bQ==",
                        },
                    }
                ],
            },
            "snippet": "6061-T6 Aluminum",
        }

        with patch(
            "bom_agent.gmail_integration._latest_inbound_message",
            return_value=inbound_message,
        ), patch(
            "bom_agent.gmail_integration.get_valid_access_token",
            return_value=(GmailCredential.objects.get(user=self.user), "access-token"),
        ), patch(
            "bom_agent.gmail_integration.parse_reply_answer",
            return_value={"answer_type": "text", "value": "6061-T6 Aluminum", "summary": "Sarah confirmed the material."},
        ), patch(
            "bom_agent.gmail_integration.maybe_resume_research",
            return_value=True,
        ) as resume_mock:
            response = self.client.post(f"/api/bom/runs/{self.run.id}/emails/poll/")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["count"], 1)
        request_obj.refresh_from_db()
        self.item.refresh_from_db()
        self.run.refresh_from_db()
        self.assertEqual(request_obj.status, "answered")
        self.assertEqual(request_obj.response, "6061-T6 Aluminum")
        self.assertIsNotNone(request_obj.answered_at)
        self.assertEqual(self.item.material_spec, "6061-T6 Aluminum")
        self.assertEqual(self.run.inputs_json[f"material_file_{self.item.id}"], "6061-T6 Aluminum")
        resume_mock.assert_called_once()
