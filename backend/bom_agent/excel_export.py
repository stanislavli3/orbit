import io
import json
import re
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from files_api.s3_service import upload_bytes_to_s3


HEADER_FILL = PatternFill("solid", fgColor="FF1F4E78")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFFFF")
BODY_FONT = Font(name="Calibri", size=11, color="FF000000")
BLUE_FONT = Font(name="Calibri", size=11, color="FF0000FF")
TITLE_FONT = Font(name="Calibri", size=11, bold=True, color="FF1F1F1F")
THIN_BORDER = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)
GREEN_FILL = PatternFill("solid", fgColor="FFC6EFCE")
AMBER_FILL = PatternFill("solid", fgColor="FFFCE4D6")
FLAG_FILL = PatternFill("solid", fgColor="FFFDE9D9")


def _stringify(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (list, dict)):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


def _safe_decimal(value) -> Decimal | None:
    if value in (None, ""):
        return None
    return Decimal(str(value))


def _best_quote(item):
    quotes = list(item.quotes.all())
    if not quotes:
        return None
    return min(
        quotes,
        key=lambda quote: (
            _safe_decimal(quote.landed_cost_usd)
            if quote.landed_cost_usd is not None
            else _safe_decimal(quote.unit_price),
            _safe_decimal(quote.unit_price),
            quote.lead_time_days,
        ),
    )


def _extract_country(notes: str) -> str:
    match = re.search(r"\[Country:\s*([^\]]+)\]", notes or "")
    return match.group(1).strip() if match else ""


def _extract_first_number(value) -> float | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, Decimal):
        return float(value)

    text = str(value)
    match = re.search(r"(\d+(?:\.\d+)?)", text.replace(",", ""))
    return float(match.group(1)) if match else None


def _timeline_days(raw_value) -> int | None:
    if raw_value in (None, ""):
        return None

    if isinstance(raw_value, (int, float, Decimal)):
        return int(raw_value)

    text = str(raw_value).strip().lower()
    match = re.search(r"(\d+(?:\.\d+)?)", text)
    if not match:
        return None

    value = float(match.group(1))
    if "week" in text:
        return int(value * 7)
    if "month" in text:
        return int(value * 30)
    return int(value)


def _item_confidence(item, results_section: dict) -> float:
    quotes = list(item.quotes.all())
    source_count = sum(1 for quote in quotes if quote.source_url)
    score = 0.3
    if item.material_spec or results_section.get("material"):
        score += 0.2
    if results_section.get("summary"):
        score += 0.15
    score += min(0.25, len(quotes) * 0.08)
    score += min(0.1, source_count * 0.05)
    return round(min(score, 0.95), 2)


def _flag_text(flag: bool) -> str:
    return "Flagged" if flag else "OK"


def _set_widths(sheet, widths: dict[str, int]) -> None:
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width


def _style_sheet(sheet) -> None:
    for row in sheet.iter_rows():
        for cell in row:
            cell.border = THIN_BORDER
            if cell.row == 1:
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                if cell.font.color is None and not cell.font.bold:
                    cell.font = BODY_FONT
                cell.alignment = Alignment(vertical="top", wrap_text=True)


def _append_summary_sheet(workbook, run, items) -> None:
    sheet = workbook.active
    sheet.title = "BOM Summary"
    headers = [
        "#",
        "Part Name",
        "Part Number",
        "Material",
        "Qty",
        "Unit Cost (Best)",
        "Extended Cost",
        "Supplier",
        "Lead Time",
        "Status",
    ]
    sheet.append(headers)

    for idx, item in enumerate(items, start=1):
        quote = _best_quote(item)
        row = sheet.max_row + 1
        sheet.append(
            [
                idx,
                item.part_name,
                item.part_number,
                item.material_spec,
                item.quantity,
                float(quote.unit_price) if quote else None,
                f"=E{row}*F{row}",
                quote.supplier_name if quote else "",
                quote.lead_time_days if quote else "",
                item.status,
            ]
        )

    total_row = max(2, sheet.max_row + 1)
    sheet[f"F{total_row}"] = "Total"
    sheet[f"F{total_row}"].font = TITLE_FONT
    end_row = max(2, total_row - 1)
    sheet[f"G{total_row}"] = f"=SUM(G2:G{end_row})"
    sheet[f"G{total_row}"].number_format = '$#,##0.00'
    sheet.freeze_panes = "A2"

    for row in sheet.iter_rows(min_row=2, max_row=total_row, min_col=6, max_col=7):
        for cell in row:
            cell.number_format = '$#,##0.00'

    _set_widths(
        sheet,
        {
            "A": 6,
            "B": 28,
            "C": 18,
            "D": 24,
            "E": 10,
            "F": 16,
            "G": 16,
            "H": 22,
            "I": 14,
            "J": 16,
        },
    )
    _style_sheet(sheet)


def _append_supplier_sheet(workbook, run, items) -> None:
    sheet = workbook.create_sheet("Supplier Comparison")
    headers = [
        "Part Name",
        "Part Number",
        "Supplier",
        "Unit Price",
        "MOQ",
        "Lead Time",
        "Tooling/NRE",
        "Certifications",
        "Notes",
        "Source URL",
    ]
    sheet.append(headers)

    for item in items:
        item_quotes = sorted(
            item.quotes.all(),
            key=lambda quote: (_safe_decimal(quote.unit_price), quote.lead_time_days),
        )
        if not item_quotes:
            sheet.append([item.part_name, item.part_number, "No quotes captured", None, None, None, None, "Not captured", "", ""])
            continue

        compliance_inputs = _stringify(run.inputs_json.get("compliance", ""))
        certifications = compliance_inputs or "Not captured"
        for quote in item_quotes:
            sheet.append(
                [
                    item.part_name,
                    item.part_number,
                    quote.supplier_name,
                    float(quote.unit_price),
                    quote.moq,
                    quote.lead_time_days,
                    float(quote.tooling_cost),
                    certifications,
                    quote.notes,
                    quote.source_url,
                ]
            )

    sheet.freeze_panes = "A2"
    max_row = sheet.max_row
    if max_row >= 2:
        best_price_formula = (
            f'$D2=MINIFS($D$2:$D${max_row},$A$2:$A${max_row},$A2,$B$2:$B${max_row},$B2)'
        )
        longest_lead_formula = (
            f'$F2=MAXIFS($F$2:$F${max_row},$A$2:$A${max_row},$A2,$B$2:$B${max_row},$B2)'
        )
        sheet.conditional_formatting.add(
            f"D2:D{max_row}",
            FormulaRule(formula=[best_price_formula], fill=GREEN_FILL),
        )
        sheet.conditional_formatting.add(
            f"F2:F{max_row}",
            FormulaRule(formula=[longest_lead_formula], fill=AMBER_FILL),
        )

    for row in sheet.iter_rows(min_row=2, max_row=max_row, min_col=4, max_col=4):
        for cell in row:
            cell.number_format = '$#,##0.00'
    for row in sheet.iter_rows(min_row=2, max_row=max_row, min_col=7, max_col=7):
        for cell in row:
            cell.number_format = '$#,##0.00'

    _set_widths(
        sheet,
        {
            "A": 28,
            "B": 18,
            "C": 24,
            "D": 14,
            "E": 10,
            "F": 12,
            "G": 14,
            "H": 22,
            "I": 40,
            "J": 40,
        },
    )
    _style_sheet(sheet)


def _append_inputs_sheet(workbook, run) -> None:
    sheet = workbook.create_sheet("Assumptions & Inputs")
    headers = ["Section", "Input", "Entered / Assumed Value", "Derived / Formula", "Notes"]
    sheet.append(headers)

    tracked_keys = [
        ("production_volume", run.inputs_json.get("production_volume", "")),
        ("timeline", run.inputs_json.get("timeline", "")),
        ("budget", run.inputs_json.get("budget", "")),
        ("exchange_rates", run.inputs_json.get("exchange_rates", run.inputs_json.get("exchange_rate", ""))),
        ("shipping_assumptions", run.inputs_json.get("shipping_assumptions", "")),
    ]
    seen = {key for key, _ in tracked_keys}
    for key, value in sorted(run.inputs_json.items()):
        if key not in seen:
            tracked_keys.append((key, value))

    row_lookup = {}
    for key, value in tracked_keys:
        row = sheet.max_row + 1
        row_lookup[key] = row
        sheet.append(["User Input", key.replace("_", " ").title(), _stringify(value), "", ""])

    standards = [
        ("Industry Standard", "Domestic handling uplift", "1.08", "", "Hardcoded default multiplier"),
        ("Industry Standard", "International shipping uplift", "1.15", "", "Hardcoded default multiplier"),
        ("Industry Standard", "FX fallback", "1.00", "", "Assume USD baseline if rate not provided"),
        ("Industry Standard", "Tooling amortization batches", "100", "", "Used for worksheet context only"),
    ]
    for standard in standards:
        sheet.append(list(standard))
        current_row = sheet.max_row
        for cell in sheet[current_row]:
            cell.font = BLUE_FONT

    production_row = row_lookup.get("production_volume")
    timeline_row = row_lookup.get("timeline")
    budget_row = row_lookup.get("budget")
    exchange_row = row_lookup.get("exchange_rates")
    shipping_row = row_lookup.get("shipping_assumptions")

    derived_rows = [
        (
            "Derived",
            "Volume upper bound",
            "",
            (
                f'=IF(C{production_row}="1–10 units",10,'
                f'IF(C{production_row}="50–100 units",100,'
                f'IF(C{production_row}="500–1 000 units",1000,'
                f'IF(C{production_row}="1 000+ units",1000,""))))'
                if production_row
                else ""
            ),
            "Formula-driven interpretation of the selected production band",
        ),
        (
            "Derived",
            "Timeline days",
            "",
            f'=IFERROR(VALUE(LEFT(C{timeline_row},FIND(" ",C{timeline_row}&" ")-1)), "")' if timeline_row else "",
            "Formula extracts the first numeric timeline value when provided",
        ),
        (
            "Derived",
            "Budget numeric",
            "",
            f'=IFERROR(VALUE(SUBSTITUTE(SUBSTITUTE(C{budget_row},"$",""),",","")), "")' if budget_row else "",
            "Formula removes currency symbols and commas",
        ),
        (
            "Derived",
            "FX multiplier in use",
            "",
            f'=IF(C{exchange_row}="",1,C{exchange_row})' if exchange_row else '=1',
            "Defaults to 1 when no exchange rate was supplied",
        ),
        (
            "Derived",
            "Shipping assumption in use",
            "",
            f'=IF(C{shipping_row}="",1.15,C{shipping_row})' if shipping_row else '=1.15',
            "Falls back to the hardcoded international shipping uplift",
        ),
    ]
    for derived in derived_rows:
        sheet.append(list(derived))
        current_row = sheet.max_row
        sheet[f"D{current_row}"].font = BODY_FONT

    sheet.freeze_panes = "A2"
    _set_widths(sheet, {"A": 18, "B": 28, "C": 24, "D": 24, "E": 36})
    _style_sheet(sheet)

    for row in range(2, sheet.max_row + 1):
        if sheet[f"A{row}"].value == "Industry Standard":
            for cell in sheet[row]:
                cell.font = BLUE_FONT


def _append_research_log_sheet(workbook, run, items) -> None:
    sheet = workbook.create_sheet("Research Log")
    headers = ["Timestamp", "Part", "Entry Type", "Detail", "Source URL", "Confidence", "Notes"]
    sheet.append(headers)

    item_map = {item.id: item for item in items}
    for entry in run.research_log or []:
        item_id = entry.get("item_id")
        item = item_map.get(item_id) if item_id else None
        results_section = (run.results_json or {}).get(f"item_{item_id}", {}) if item_id else {}
        confidence = _item_confidence(item, results_section) if item else ""
        sheet.append(
            [
                entry.get("ts", ""),
                item.part_name if item else entry.get("part_name", ""),
                entry.get("type", ""),
                entry.get("message", ""),
                entry.get("source_url", ""),
                confidence,
                entry.get("query", "") or entry.get("source", ""),
            ]
        )

    for item in items:
        results_section = (run.results_json or {}).get(f"item_{item.id}", {})
        confidence = _item_confidence(item, results_section)
        sheet.append(
            [
                run.completed_at.isoformat() if run.completed_at else "",
                item.part_name,
                "line_item_confidence",
                results_section.get("summary", "Research summary captured"),
                "",
                confidence,
                "Confidence inferred from material coverage, quote count, and source coverage",
            ]
        )
        for quote in item.quotes.all():
            sheet.append(
                [
                    run.completed_at.isoformat() if run.completed_at else "",
                    item.part_name,
                    "data_point",
                    f"{quote.supplier_name}: ${quote.unit_price}/unit, MOQ {quote.moq}, {quote.lead_time_days}d lead",
                    quote.source_url,
                    confidence,
                    quote.notes,
                ]
            )

    sheet.freeze_panes = "A2"
    _set_widths(sheet, {"A": 14, "B": 26, "C": 18, "D": 46, "E": 42, "F": 12, "G": 40})
    _style_sheet(sheet)


def _append_risk_sheet(workbook, run, items) -> None:
    sheet = workbook.create_sheet("Risk Matrix")
    headers = [
        "Part Name",
        "Part Number",
        "Single-Source",
        "Long Lead Time",
        "Compliance Gap",
        "Geographic Concentration",
        "Deadline Risk",
        "Summary",
    ]
    sheet.append(headers)

    deadline_days = _timeline_days(run.inputs_json.get("timeline"))

    for item in items:
        results_section = (run.results_json or {}).get(f"item_{item.id}", {})
        quotes = list(item.quotes.all())
        best_lead = min((quote.lead_time_days for quote in quotes), default=None)
        risk_flags = set(results_section.get("risk_flags", []))
        deadline_risk = bool(deadline_days and best_lead and best_lead > deadline_days)
        row = [
            item.part_name,
            item.part_number,
            _flag_text("single-source" in risk_flags),
            _flag_text("long-lead-time" in risk_flags),
            _flag_text("compliance-gap" in risk_flags),
            _flag_text("geographic-concentration" in risk_flags),
            _flag_text(deadline_risk),
            results_section.get("risk_summary", ""),
        ]
        sheet.append(row)
        current_row = sheet.max_row
        for column in ("C", "D", "E", "F", "G"):
            if sheet[f"{column}{current_row}"].value == "Flagged":
                sheet[f"{column}{current_row}"].fill = FLAG_FILL

    sheet.freeze_panes = "A2"
    _set_widths(sheet, {"A": 28, "B": 18, "C": 16, "D": 16, "E": 16, "F": 24, "G": 16, "H": 44})
    _style_sheet(sheet)


def build_bom_workbook(run):
    workbook = Workbook(write_only=False)
    workbook.remove(workbook.active)
    workbook.create_sheet("BOM Summary")

    items = list(run.line_items.prefetch_related("quotes").all())
    _append_summary_sheet(workbook, run, items)
    _append_supplier_sheet(workbook, run, items)
    _append_inputs_sheet(workbook, run)
    _append_research_log_sheet(workbook, run, items)
    _append_risk_sheet(workbook, run, items)
    return workbook


def render_bom_workbook(run) -> bytes:
    workbook = build_bom_workbook(run)
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def upload_bom_workbook(run) -> str:
    workbook_bytes = render_bom_workbook(run)
    key = f"bom-reports/{run.project_id}/run-{run.id}/bom-research-run-{run.id}.xlsx"
    upload_bytes_to_s3(
        workbook_bytes,
        key,
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    return key
